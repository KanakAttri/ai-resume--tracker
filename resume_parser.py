
from pathlib import Path
import re


def extract_text(file_path):
    """Extract text from a PDF or DOCX resume."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )
        return text

    elif extension == ".docx":
        from docx import Document

        document = Document(str(path))
        parts = []

        # Extract paragraphs
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())

        # Extract content inside tables
        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

        return "\n".join(parts)

    else:
        raise ValueError("Only PDF and DOCX files are supported.")


def extract_sections(text):
    """Detect common resume sections using their headings."""

    headings = {
        "skills": [
            "skills", "technical skills", "key skills",
            "core competencies"
        ],
        "education": [
            "education", "academic background",
            "qualifications"
        ],
        "experience": [
            "experience", "work experience",
            "professional experience", "internship"
        ],
        "projects": [
            "projects", "academic projects",
            "personal projects"
        ],
        "certifications": [
            "certifications", "certificates"
        ]
    }

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    sections = {name: "" for name in headings}
    current_section = None

    for line in lines:
        heading = re.sub(r"[:\s]+$", "", line.lower()).strip()

        found_section = None
        for section, aliases in headings.items():
            if heading in aliases:
                found_section = section
                break

        if found_section:
            current_section = found_section
            continue

        if current_section:
            sections[current_section] += line + "\n"

    return {
        section: content.strip()
        for section, content in sections.items()
    }


def parse_resume(file_path):
    """Return extracted text and structured resume information."""

    text = extract_text(file_path)
    sections = extract_sections(text)

    email_match = re.search(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text
    )

    phone_match = re.search(
        r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)",
        text
    )

    return {
        "text": text,
        "email": email_match.group(0) if email_match else "",
        "phone": phone_match.group(0).strip() if phone_match else "",
        "sections": sections
    }
